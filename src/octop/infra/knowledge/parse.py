"""Text extraction for the document types accepted by knowledge bases."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import struct
from collections.abc import Iterable
from html.parser import HTMLParser
from pathlib import Path
from typing import TYPE_CHECKING, Any

from octop.infra.knowledge.ocr import OCR_IMAGE_SUFFIXES

if TYPE_CHECKING:
    from octop.infra.knowledge.ocr import OcrExtractor

_PLAIN_TEXT_SUFFIXES = {
    ".md",
    ".markdown",
    ".txt",
    ".rst",
    ".yaml",
    ".yml",
    ".jsonl",
}
_DOCX_FALLBACK_TAG = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_DOCX_HTML_TYPES = {"application/xhtml+xml", "text/html"}
_DOCX_MAIN_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MIN_IMAGE_OCR_CHARS = 10
_MIN_CONTENT_IMAGE_PX = 80
_VECTOR_IMAGE_EXTS = {"emf", "wmf"}


def parse_document(path: Path, *, ocr: OcrExtractor | None = None) -> str:
    """Extract searchable text from a supported local document."""
    return parse_document_with_assets(path, ocr=ocr)[0]


def parse_document_with_assets(
    path: Path, *, ocr: OcrExtractor | None = None
) -> tuple[str, dict[str, bytes]]:
    """Return searchable text and the unique raster images referenced from it.

    PDF pages and PPTX slides keep each content image beside that page's text via
    ``[[IMG:{id}]]``. Office zip text (docx, xlsx) stays appended, without markers.
    A standalone image keeps its OCR text only; callers fetch it with ``self``.
    """
    suffix = path.suffix.lower()
    if suffix in OCR_IMAGE_SUFFIXES:
        if ocr is None:
            raise RuntimeError("knowledge OCR is not enabled")
        return ocr(path), {}
    if suffix in _PLAIN_TEXT_SUFFIXES:
        return _read_text(path), {}
    if suffix == ".json":
        return _parse_json(path), {}
    if suffix in {".html", ".htm"}:
        return _parse_html(path), {}
    if suffix == ".csv":
        return _parse_delimited(path, delimiter=","), {}
    if suffix == ".tsv":
        return _parse_delimited(path, delimiter="\t"), {}
    if suffix == ".pdf":
        return _parse_pdf(path, ocr)
    if suffix == ".docx":
        return _with_zip_images(path, _parse_docx(path), "word/media/", ocr), {}
    if suffix in {".wps", ".doc"}:
        from octop.infra.utils.anydoc_convert import extract_markdown_file

        text = extract_markdown_file(path)
        if not text.strip():
            raise ValueError(f"failed to extract text from {suffix} document")
        return text, {}
    if suffix == ".pptx":
        return _parse_pptx(path, ocr)
    if suffix in {".xlsx", ".xlsm"}:
        return _with_zip_images(path, _parse_xlsx(path), "xl/media/", ocr), {}
    if suffix == ".xls":
        return _parse_xls(path), {}
    raise ValueError(f"unsupported knowledge document extension: {suffix or '(none)'}")


class _ImageBag:
    """First-seen rasters for one document. Repeated bytes and icons are dropped."""

    def __init__(self) -> None:
        self.images: dict[str, bytes] = {}
        self._seen: set[bytes] = set()

    def add(
        self,
        data: bytes,
        *,
        ext: str = "",
        width: int | None = None,
        height: int | None = None,
    ) -> str | None:
        if ext.lower().lstrip(".") in _VECTOR_IMAGE_EXTS:
            return None
        if _is_icon(data, width=width, height=height):
            return None
        digest = hashlib.sha256(data).digest()
        if digest in self._seen:
            return None
        self._seen.add(digest)
        image_id = digest.hex()[:12]
        self.images[image_id] = data
        return image_id


def _marker_block(image_id: str, ocr_text: str) -> str:
    body = ocr_text.strip()
    if body:
        return f"[[IMG:{image_id}]]\n{body}"
    return f"[[IMG:{image_id}]]"


def _parse_pdf(path: Path, ocr: OcrExtractor | None) -> tuple[str, dict[str, bytes]]:
    from pypdf import PdfReader

    page_texts = [page.extract_text() or "" for page in PdfReader(path).pages]
    text = "\n".join(page_texts)
    if ocr is None or not text.strip():
        if ocr is not None and not text.strip():
            return ocr(path), {}
        return text, {}

    import pymupdf

    try:
        document = pymupdf.open(path)
    except Exception:
        return text, {}
    bag = _ImageBag()
    sections: list[str] = []
    marked = False
    try:
        for index, page_text in enumerate(page_texts):
            parts = [page_text] if page_text else []
            if index < document.page_count:
                for image in document[index].get_images(full=True):
                    try:
                        meta = document.extract_image(image[0])
                        data = meta["image"]
                    except Exception:
                        continue
                    image_id = bag.add(
                        data,
                        ext=str(meta.get("ext") or ""),
                        width=_as_int(meta.get("width")),
                        height=_as_int(meta.get("height")),
                    )
                    if image_id is None:
                        continue
                    parts.append(_marker_block(image_id, ocr(data)))
                    marked = True
            if parts:
                sections.append("\n".join(parts))
    finally:
        document.close()
    if not marked:
        return text, {}
    return "\n".join(sections), bag.images


def _parse_pptx(path: Path, ocr: OcrExtractor | None) -> tuple[str, dict[str, bytes]]:
    from pptx import Presentation

    bag = _ImageBag()
    pieces: list[str] = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                pieces.append(shape.text)
            if ocr is None:
                continue
            try:
                image = getattr(shape, "image", None)
                if image is None:
                    continue
                data = image.blob
                ext = str(getattr(image, "ext", "") or "")
            except Exception:
                continue
            image_id = bag.add(data, ext=ext)
            if image_id is None:
                continue
            pieces.append(_marker_block(image_id, ocr(data)))
    return "\n".join(pieces), bag.images if ocr is not None else {}


def _as_int(value: object) -> int | None:
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return int(value)


def _is_icon(data: bytes, *, width: int | None = None, height: int | None = None) -> bool:
    if width is not None and height is not None:
        return width < _MIN_CONTENT_IMAGE_PX and height < _MIN_CONTENT_IMAGE_PX
    size = _raster_size(data)
    if size is None:
        return False
    return size[0] < _MIN_CONTENT_IMAGE_PX and size[1] < _MIN_CONTENT_IMAGE_PX


def _raster_size(data: bytes) -> tuple[int, int] | None:
    if data.startswith(b"\x89PNG\r\n\x1a\n") and len(data) >= 24:
        return struct.unpack(">II", data[16:24])
    if data.startswith((b"GIF87a", b"GIF89a")) and len(data) >= 10:
        return struct.unpack("<HH", data[6:10])
    if data.startswith(b"\xff\xd8"):
        return _jpeg_size(data)
    if len(data) >= 30 and data.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return _webp_size(data)
    return None


def _jpeg_size(data: bytes) -> tuple[int, int] | None:
    index = 2
    limit = len(data)
    while index + 8 < limit:
        if data[index] != 0xFF:
            index += 1
            continue
        marker = data[index + 1]
        if marker in {0xC0, 0xC1, 0xC2}:
            height, width = struct.unpack(">HH", data[index + 5 : index + 9])
            return width, height
        if marker in {0xD8, 0xD9}:
            break
        if index + 4 > limit:
            break
        segment = struct.unpack(">H", data[index + 2 : index + 4])[0]
        if segment < 2:
            break
        index += 2 + segment
    return None


def _webp_size(data: bytes) -> tuple[int, int] | None:
    chunk = data[12:16]
    if chunk == b"VP8X" and len(data) >= 30:
        width = 1 + int.from_bytes(data[24:27], "little")
        height = 1 + int.from_bytes(data[27:30], "little")
        return width, height
    if chunk == b"VP8 " and len(data) >= 30:
        width, height = struct.unpack("<HH", data[26:30])
        return width & 0x3FFF, height & 0x3FFF
    if chunk == b"VP8L" and len(data) >= 25:
        bits = int.from_bytes(data[21:25], "little")
        return (bits & 0x3FFF) + 1, ((bits >> 14) & 0x3FFF) + 1
    return None


_ZIP_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}


def _with_zip_images(path: Path, text: str, media_prefix: str, ocr: OcrExtractor | None) -> str:
    """Keep extracted text, and append OCR of pictures packed in an Office zip."""
    if ocr is None:
        return text
    extra = _zip_media_images_text(path, media_prefix, ocr)
    if not extra:
        return text
    return f"{text}\n{extra}" if text.strip() else extra


def _zip_media_images_text(path: Path, media_prefix: str, ocr: OcrExtractor) -> str:
    import hashlib
    import zipfile

    pieces: list[str] = []
    seen: set[bytes] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith(media_prefix) or name.endswith("/"):
                continue
            if Path(name).suffix.lower() not in _ZIP_IMAGE_SUFFIXES:
                continue
            data = archive.read(name)
            digest = hashlib.sha256(data).digest()
            if digest in seen:
                continue
            seen.add(digest)
            text = ocr(data)
            if len(text.strip()) >= _MIN_IMAGE_OCR_CHARS:
                pieces.append(text)
    return "\n".join(pieces)


_TEXT_ENCODINGS = ("utf-8-sig", "gb18030")


def _read_text(path: Path) -> str:
    """Decode a knowledge text file, preferring UTF-8 then GB18030 (GBK).

    ``errors='replace'`` on UTF-8 used to swallow GBK Windows documents as
    mojibake, so retrieval never saw the original Chinese text.
    """
    data = path.read_bytes()
    for encoding in _TEXT_ENCODINGS:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        else:
            return text.replace("\r\n", "\n").replace("\r", "\n")
    return data.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


def _parse_json(path: Path) -> str:
    raw = _read_text(path)
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    return json.dumps(parsed, ensure_ascii=False, indent=2)


def _parse_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return _docx_body_text(document.element.body, document.part)


def _docx_body_text(body: Any, part: Any) -> str:
    """Collect every ``w:p`` and ``w:altChunk`` in a body, in document order.

    ``Document.paragraphs`` only lists ``w:p`` children of ``w:body``, so text inside
    tables, content controls (``w:sdt``), revision wrappers (``w:ins``), and text boxes
    is silently dropped. ``mc:Fallback`` duplicates its ``mc:Choice`` sibling, so its
    paragraphs are skipped.
    """
    from docx.oxml.ns import qn

    alt_chunk_tag = qn("w:altChunk")
    lines: list[str] = []
    for element in body.iter(qn("w:p"), alt_chunk_tag):
        if next(element.iterancestors(_DOCX_FALLBACK_TAG), None) is not None:
            continue
        if element.tag == alt_chunk_tag:
            lines.append(_docx_alt_chunk_text(element, part))
        else:
            lines.append(element.text)
    return "\n".join(lines)


def _docx_alt_chunk_text(element: Any, part: Any) -> str:
    """Expand a ``w:altChunk`` reference the way Word does when opening the file.

    Converters keep the bulk of the text in an embedded HTML or Word part and leave
    only a stub in ``document.xml``; unexpanded, that body is lost.
    """
    from docx.oxml.ns import qn

    chunk = part.related_parts.get(element.get(qn("r:id")) or "")
    if chunk is None:
        return ""
    content_type = str(chunk.content_type)
    blob: bytes = chunk.blob
    if content_type in _DOCX_HTML_TYPES:
        return _html_text(blob.decode("utf-8-sig", errors="replace"))
    if content_type == "text/plain":
        return blob.decode("utf-8-sig", errors="replace")
    if content_type == _DOCX_MAIN_TYPE:
        from docx import Document

        nested = Document(io.BytesIO(blob))
        return _docx_body_text(nested.element.body, nested.part)
    return ""


def _parse_html(path: Path) -> str:
    return _html_text(_read_text(path))


def _html_text(raw: str) -> str:
    parser = _HTMLTextParser()
    parser.feed(raw)
    parser.close()
    lines = [" ".join(line.split()) for line in parser.text().splitlines()]
    return "\n".join(line for line in lines if line)


class _HTMLTextParser(HTMLParser):
    _SKIP = frozenset({"script", "style", "noscript", "template"})
    _BLOCK = frozenset(
        {
            "p",
            "div",
            "br",
            "li",
            "tr",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "article",
            "section",
            "header",
            "footer",
            "blockquote",
            "pre",
            "table",
            "ul",
            "ol",
        }
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, _attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip += 1
            return
        if not self._skip and tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            if self._skip:
                self._skip -= 1
            return
        if not self._skip and tag in self._BLOCK:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip and data:
            self._parts.append(data)

    def text(self) -> str:
        return "".join(self._parts)


def _stringify_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", " ").replace("\n", " ").strip()


def _trim_row(cells: list[str]) -> list[str]:
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def _sheet_text(title: str, rows: Iterable[Iterable[object]]) -> str:
    lines = [f"# {title}"]
    has_cells = False
    for raw in rows:
        cells = _trim_row([_stringify_cell(cell) for cell in raw])
        if not cells:
            continue
        has_cells = True
        lines.append("\t".join(cells))
    if not has_cells:
        return ""
    return "\n".join(lines)


def _parse_delimited(path: Path, *, delimiter: str) -> str:
    reader = csv.reader(_read_text(path).splitlines(), delimiter=delimiter)
    return _sheet_text(path.stem, reader)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            text = _sheet_text(str(sheet.title), sheet.iter_rows(values_only=True))
            if text:
                parts.append(text)
        return "\n\n".join(parts)
    finally:
        workbook.close()


def _parse_xls(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=False)
    parts: list[str] = []
    for sheet in book.sheets():
        rows: list[list[object]] = []
        for row_idx in range(sheet.nrows):
            rows.append(
                [
                    _xlrd_cell_value(book, sheet.cell(row_idx, col_idx))
                    for col_idx in range(sheet.ncols)
                ]
            )
        text = _sheet_text(sheet.name, rows)
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _xlrd_cell_value(book: object, cell: object) -> object:
    import xlrd

    ctype = getattr(cell, "ctype", xlrd.XL_CELL_EMPTY)
    value = getattr(cell, "value", None)
    if ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return None
    if ctype == xlrd.XL_CELL_DATE:
        datemode = int(getattr(book, "datemode", 0))
        try:
            return xlrd.xldate_as_datetime(value, datemode).isoformat(sep=" ", timespec="seconds")
        except (OSError, OverflowError, TypeError, ValueError):
            return value
    return value
